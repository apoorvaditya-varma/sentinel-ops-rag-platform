import io
import zipfile

import pandas as pd
import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from src.guardrails.file_validation import FileRejected, safe_filename, validate_upload


def make_pdf(text: str = "Patch critical CVEs within 7 days.") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


def make_docx(text: str = "Runbook: restart the ingress controller.") -> bytes:
    buffer = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("Runbook", level=1)
    doc.add_paragraph(text)
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    buffer = io.BytesIO()
    pd.DataFrame({"cve_id": ["CVE-2024-6387"], "severity": ["High"]}).to_excel(buffer, index=False)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "name, data",
    [
        ("policy.pdf", make_pdf()),
        ("runbook.docx", make_docx()),
        ("scan.xlsx", make_xlsx()),
        ("alerts.csv", b"alert_id,severity\nA-1,High\n"),
        ("notes.md", b"# Runbook\nkubectl rollout undo deploy/api"),
        ("notes.txt", "Caf\u00e9 VPN guide".encode("cp1252")),
        ("iam.json", b'{"Effect": "Allow", "Action": "s3:GetObject"}'),
        ("events.json", b'{"a": 1}\n{"a": 2}\n'),
        ("deploy.yaml", b"apiVersion: apps/v1\nkind: Deployment\n---\nkind: Service\n"),
        ("app.log", b"2026-10-14 09:12:01 ERROR pool exhausted\n"),
    ],
)
def test_valid_files_accepted(name, data):
    result = validate_upload(name, data)
    assert result.filename == name
    assert len(result.sha256) == 64


@pytest.mark.parametrize(
    "name, data, reason",
    [
        ("old.doc", b"\xd0\xcf\x11\xe0", "Legacy Word"),
        ("old.xls", b"\xd0\xcf\x11\xe0", "Legacy Excel"),
        ("tool.exe", b"MZ\x90\x00", "Executables"),
        ("deploy.sh", b"#!/bin/bash\nrm -rf /", "Executables"),
        ("bundle.zip", b"PK\x03\x04", "Archives"),
        ("diagram.png", b"\x89PNG", "Images"),
        ("macro.docm", b"PK\x03\x04", "Macro-enabled"),
        ("unknown.xyz", b"hello", "not accepted"),
        ("fake.pdf", b"this is not a pdf at all", "not a PDF"),
        ("fake.docx", b"plain text pretending", "not an Office document"),
        ("invoice.exe.pdf", make_pdf(), "double extension"),
        ("empty.txt", b"", "empty"),
        ("binary.txt", b"abc\x00\x01\x02", "binary"),
        ("bad.json", b"{not json", "Invalid JSON"),
        ("bad.yaml", b"key: [unclosed", "Invalid YAML"),
        ("one_col.csv", b"just_one_column\nvalue\n", "two columns"),
        ("blank.pdf", make_pdf(" "), "no extractable text"),
    ],
)
def test_invalid_files_rejected(name, data, reason):
    with pytest.raises(FileRejected, match=reason):
        validate_upload(name, data)


def test_macro_inside_docx_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<w:document/>")
        archive.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(FileRejected, match="macros"):
        validate_upload("sneaky.docx", buffer.getvalue())


def test_oversized_file_rejected(monkeypatch):
    from dataclasses import replace

    from src.guardrails import file_validation

    small = replace(file_validation.get_settings(), max_upload_mb=1)
    monkeypatch.setattr(file_validation, "get_settings", lambda: small)
    with pytest.raises(FileRejected, match="limit is 1 MB"):
        validate_upload("big.txt", b"a" * (2 * 1024 * 1024))


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("../../etc/passwd.txt", "passwd.txt"),
        ("C:\\Users\\me\\report.pdf", "report.pdf"),
        ("rapport d'audit\u00a0(v2).pdf", "rapport d_audit _v2_.pdf"),
    ],
)
def test_safe_filename(raw, expected):
    assert safe_filename(raw) == expected


def test_hidden_filename_rejected():
    with pytest.raises(FileRejected):
        safe_filename(".env")
