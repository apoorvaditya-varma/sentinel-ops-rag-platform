"""Task 3: file loaders -> LangChain Documents with source metadata.

Every loader returns a list of Documents whose metadata always contains
`source` (file name) and `file_type`; PDFs add `page`, tables add `sheet`
and `rows`.  Tabular files are also exposed as DataFrames for DuckDB.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from src.guardrails.file_validation import decode_text, read_csv_text

ROWS_PER_DOCUMENT = 25


def load_file(path: Path) -> list[Document]:
    ext = path.suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"No loader for {ext}")
    docs = [doc for doc in loader(path) if doc.page_content.strip()]
    for doc in docs:
        doc.metadata = {"source": path.name, "file_type": ext.lstrip("."), **doc.metadata}
    return docs


def load_tables(path: Path) -> dict[str, pd.DataFrame]:
    """Return {table_name: DataFrame} for a CSV (one table) or Excel file (one per sheet)."""
    if path.suffix.lower() == ".csv":
        return {path.stem: read_csv_text(decode_text(path.read_bytes()))}
    sheets = pd.read_excel(path, sheet_name=None)
    return {f"{path.stem}__{name}": frame for name, frame in sheets.items() if not frame.empty}


def _load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(path)
    return [
        Document(page_content=page.extract_text() or "", metadata={"page": number})
        for number, page in enumerate(reader.pages, start=1)
    ]


def _load_docx(path: Path) -> list[Document]:
    doc = DocxDocument(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") and style[-1:].isdigit():
            text = f"{'#' * int(style[-1])} {text}"
        parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Document(page_content="\n".join(parts))]


def _load_text(path: Path) -> list[Document]:
    return [Document(page_content=decode_text(path.read_bytes()))]


def _load_json(path: Path) -> list[Document]:
    text = decode_text(path.read_bytes())
    try:
        text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        pass  # JSON Lines: keep one record per line
    return [Document(page_content=text)]


def _load_tabular(path: Path) -> list[Document]:
    docs: list[Document] = []
    for name, frame in load_tables(path).items():
        frame = frame.dropna(how="all")
        sheet = name.split("__", 1)[1] if "__" in name else None
        columns = ", ".join(map(str, frame.columns))
        summary = f"Table '{name}' with {len(frame)} rows and columns: {columns}."
        docs.append(Document(page_content=summary, metadata={"sheet": sheet, "rows": "schema"}))
        for start in range(0, len(frame), ROWS_PER_DOCUMENT):
            block = frame.iloc[start : start + ROWS_PER_DOCUMENT]
            lines = [
                "; ".join(f"{col}: {val}" for col, val in row.items() if pd.notna(val))
                for _, row in block.iterrows()
            ]
            docs.append(
                Document(
                    page_content=f"Table '{name}' rows {start + 1}-{start + len(block)}:\n" + "\n".join(lines),
                    metadata={"sheet": sheet, "rows": f"{start + 1}-{start + len(block)}"},
                )
            )
    return docs


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_text,
    ".txt": _load_text,
    ".log": _load_text,
    ".yaml": _load_text,
    ".yml": _load_text,
    ".json": _load_json,
    ".csv": _load_tabular,
    ".xlsx": _load_tabular,
}


def text_windows(docs: list[Document], count: int = 6, size: int = 800) -> list[str]:
    """Evenly spaced excerpts across the whole file, used for screening."""
    text = "\n".join(doc.page_content for doc in docs).strip()
    if len(text) <= size * count:
        return [text[i : i + size] for i in range(0, len(text), size)] or [""]
    step = (len(text) - size) // (count - 1)
    return [text[i * step : i * step + size] for i in range(count)]
